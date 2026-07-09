"""
Bulk re-ingestion script.

Uses the Excel reference sheet as the source of truth for employee IDs and
filenames, then resolves each filename to its actual path by searching the
known resume folders. This covers all 252 uploaded employees regardless of
what resume_store has.

Usage:
    python db/reingest_resumes.py [--limit N] [--employee-id SS519]
    python db/reingest_resumes.py --concurrency 3

Options:
    --limit N           Only process the first N resumes (useful for testing)
    --employee-id ID    Only re-ingest a single employee
    --concurrency N     Max parallel Gemini calls (default: 3)
    --dry-run           Print what would be processed without calling Gemini
"""

import asyncio
import argparse
import logging
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from tools.resume_tool import extract_resume

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)
logger = logging.getLogger(__name__)

EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "excel_data", "Employee_Master_Reference_20260707_164120.xlsx"
)

# All folders to search for resume files (order matters — first match wins)
SEARCH_ROOTS = [
    r"C:\Users\LENOVO\OneDrive - Sails Software Solutions Pvt Ltd\resumes_sails\Sails Resumes",
    r"C:\Users\LENOVO\Downloads\resume-sync\backend\resumes",
    r"C:\Users\LENOVO\Downloads\resume-sync\backend\mock_resumes_data",
    r"C:\Users\LENOVO\Downloads\resume-sync\backend\Sails Resumes",
]


def build_file_map() -> dict[str, str]:
    """Walk all search roots and return {lowercase_filename: full_path}."""
    file_map = {}
    for root in SEARCH_ROOTS:
        if not os.path.exists(root):
            continue
        for dirpath, _, filenames in os.walk(root):
            for fn in filenames:
                if fn.lower().endswith('.docx'):
                    key = fn.lower()
                    if key not in file_map:  # first match wins
                        file_map[key] = os.path.join(dirpath, fn)
    return file_map


def load_excel_entries() -> list[dict]:
    """Read all Uploaded rows from the Excel master sheet."""
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Master Reference']
    headers = [cell.value for cell in ws[2]]
    entries = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        d = dict(zip(headers, row))
        if d.get('Resume Status') == 'Uploaded' and d.get('Resume Filename'):
            entries.append({
                'employee_id': d['Employee ID'],
                'name': d['Employee Name'],
                'filename': d['Resume Filename'],
            })
    return entries


async def reingest_one(employee_id: str, file_path: str, semaphore: asyncio.Semaphore) -> dict:
    async with semaphore:
        logger.info(f"[{employee_id}] {os.path.basename(file_path)}")
        try:
            result = await extract_resume(employee_id=employee_id, file_path=file_path)
            if result.get("status") == "error":
                logger.error(f"[{employee_id}] FAILED: {result.get('message')}")
                return {"employee_id": employee_id, "status": "error", "message": result.get("message")}
            logger.info(f"[{employee_id}] OK")
            return {"employee_id": employee_id, "status": "success"}
        except Exception as e:
            logger.error(f"[{employee_id}] EXCEPTION: {e}")
            return {"employee_id": employee_id, "status": "error", "message": str(e)}


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--employee-id", type=str, default=None)
    parser.add_argument("--concurrency", type=int, default=3)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    logger.info("Building file map from resume folders...")
    file_map = build_file_map()
    logger.info(f"Found {len(file_map)} .docx files across search roots")

    logger.info("Loading employee list from Excel...")
    entries = load_excel_entries()
    logger.info(f"Excel has {len(entries)} uploaded employees")

    # Filter to single employee if requested
    if args.employee_id:
        entries = [e for e in entries if e['employee_id'] == args.employee_id]

    # Resolve filenames to full paths
    to_process = []
    not_found = []
    for e in entries:
        key = e['filename'].lower()
        path = file_map.get(key)
        if path:
            to_process.append({'employee_id': e['employee_id'], 'name': e['name'], 'file_path': path})
        else:
            not_found.append(e)

    logger.info(f"Resolved: {len(to_process)} | Not found on disk: {len(not_found)}")
    if not_found:
        logger.warning("Files not found on disk:")
        for e in not_found:
            logger.warning(f"  [{e['employee_id']}] {e['filename']}")

    if args.limit:
        to_process = to_process[:args.limit]
        logger.info(f"Limiting to {args.limit} resumes")

    if args.dry_run:
        logger.info("DRY RUN — would process:")
        for d in to_process:
            logger.info(f"  [{d['employee_id']}] {d['name']} — {os.path.basename(d['file_path'])}")
        return

    if not to_process:
        logger.info("Nothing to process.")
        return

    logger.info(f"Starting re-ingestion of {len(to_process)} resumes (concurrency={args.concurrency})")

    semaphore = asyncio.Semaphore(args.concurrency)
    tasks = [reingest_one(d['employee_id'], d['file_path'], semaphore) for d in to_process]
    results = await asyncio.gather(*tasks)

    succeeded = [r for r in results if r["status"] == "success"]
    failed = [r for r in results if r["status"] == "error"]

    print("\n" + "=" * 60)
    print("Re-ingestion complete")
    print(f"  Success      : {len(succeeded)}")
    print(f"  Failed       : {len(failed)}")
    print(f"  Not on disk  : {len(not_found)}")
    if failed:
        print("\nFailed employees:")
        for f in failed:
            print(f"  [{f['employee_id']}] {f.get('message', '')}")
    print("=" * 60)


if __name__ == "__main__":
    asyncio.run(main())
